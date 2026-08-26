"""
llm_analyzer.py — Hybrid LLM + Code-Execution Excel Q&A Engine.

Uses Google Gemini (FREE) by default. Set GEMINI_API_KEY to use.
Optionally supports Anthropic Claude — set ANTHROPIC_API_KEY instead.

Pipeline:
  1. LLM reads DataFrame schema → generates pandas code
  2. Python executes code on FULL dataset → exact numbers
  3. LLM formats raw output → clean readable text
"""

import io
import json
import os
import traceback
import logging
from typing import Any

import pandas as pd
import numpy as np

logger = logging.getLogger(__name__)
MAX_RETRIES = 2


# ═══════════════════════════════════════════════════════════════════════════
# LLM CLIENT — auto-detects Gemini (free) or Anthropic (paid)
# ═══════════════════════════════════════════════════════════════════════════

def _get_provider() -> str:
    """Detect which API key is set."""
    if os.getenv("GEMINI_API_KEY"):
        return "gemini"
    if os.getenv("ANTHROPIC_API_KEY"):
        return "anthropic"
    raise EnvironmentError(
        "No API key found. Set GEMINI_API_KEY (free) or ANTHROPIC_API_KEY."
    )


def _call_llm(system: str, user_message: str, model: str | None = None) -> str:
    """Call the LLM (Gemini or Anthropic) and return text response."""
    provider = _get_provider()

    if provider == "gemini":
        import google.generativeai as genai
        genai.configure(api_key=os.environ["GEMINI_API_KEY"])
        mdl = model or "gemini-3.6-flash"
        gm = genai.GenerativeModel(mdl, system_instruction=system)
        resp = gm.generate_content(user_message)
        return resp.text

    else:  # anthropic
        from anthropic import Anthropic
        client = Anthropic(api_key=os.environ["ANTHROPIC_API_KEY"])
        mdl = model or "claude-sonnet-4-20250514"
        resp = client.messages.create(
            model=mdl, max_tokens=4096, system=system,
            messages=[{"role": "user", "content": user_message}],
        )
        return resp.content[0].text


# ═══════════════════════════════════════════════════════════════════════════
# DATAFRAME LOADER — loads FULL data for code execution
# ═══════════════════════════════════════════════════════════════════════════

def load_dataframes(file_path: str, file_name: str) -> dict[str, pd.DataFrame]:
    """Load every sheet into DataFrame(s). A sheet holding several stacked
    tables (separated by blank rows — common in dashboard-style reports) is
    split into one DataFrame per table so unrelated tables never get mixed."""
    from pathlib import Path
    ext = Path(file_name).suffix.lower()

    if ext in (".csv", ".tsv"):
        sep = "\t" if ext == ".tsv" else ","
        return {"Sheet1": pd.read_csv(file_path, sep=sep, on_bad_lines="skip")}

    from openpyxl import load_workbook
    wb = load_workbook(file_path, data_only=True, read_only=False)
    dfs: dict[str, pd.DataFrame] = {}

    for ws_name in wb.sheetnames:
        ws = wb[ws_name]
        if not ws.max_row or ws.max_row == 0:
            continue

        # Unmerge + fill
        fill: dict[tuple[int,int], Any] = {}
        for mr in list(ws.merged_cells.ranges):
            anchor = ws.cell(mr.min_row, mr.min_col).value
            for r in range(mr.min_row, mr.max_row+1):
                for c in range(mr.min_col, mr.max_col+1):
                    fill[(r,c)] = anchor

        rows = []
        for r in range(1, ws.max_row+1):
            rows.append([fill.get((r,c), ws.cell(r,c).value) for c in range(1, (ws.max_column or 0)+1)])

        tables = _split_tables(rows)
        if not tables:
            continue

        for title, header_row, data in tables:
            headers = _dedup([str(v) if v is not None else f"Col{j+1}" for j,v in enumerate(header_row)])
            df = pd.DataFrame(data, columns=headers)

            # Auto-convert types
            for col in df.columns:
                num = pd.to_numeric(df[col], errors="coerce")
                if num.notna().sum() > df[col].notna().sum() * 0.5:
                    df[col] = num
                    continue
                try:
                    dt = pd.to_datetime(df[col], errors="coerce", infer_datetime_format=True)
                    if dt.notna().sum() > df[col].notna().sum() * 0.5:
                        df[col] = dt
                except Exception:
                    pass

            df = df.dropna(how="all").reset_index(drop=True)
            df = df.ffill()
            if df.empty:
                continue

            if len(tables) == 1:
                name = ws_name
            elif title:
                short = title.split("|")[0].strip()[:60]
                name = f"{ws_name} - {short}" if short else ws_name
            else:
                name = f"{ws_name} (Table {len(dfs)+1})"
            base, i = name, 2
            while name in dfs:
                name = f"{base} ({i})"; i += 1
            dfs[name] = df

    wb.close()
    return dfs


def _split_tables(rows: list[list]) -> list[tuple[str | None, list, list]]:
    """Split sheet rows into separate tables at fully-blank rows.

    Returns a list of (title, header_row, data_rows) — one per detected table.
    A block with no data rows below its header candidate is treated as a
    standalone title line and attached to the next real table instead.
    """
    blocks: list[list] = []
    cur: list = []
    for row in rows:
        if all(v is None for v in row):
            if cur:
                blocks.append(cur)
            cur = []
        else:
            cur.append(row)
    if cur:
        blocks.append(cur)

    tables: list[tuple[str | None, list, list]] = []
    pending_title = None
    for block in blocks:
        h_idx = _find_header_row(block)
        header_row = block[h_idx]
        data_rows = block[h_idx+1:]
        if not data_rows:
            text = " ".join(str(v) for v in block[-1] if v is not None)
            pending_title = text or pending_title
            continue
        header_row, data_rows = _trim_empty_cols(header_row, data_rows)
        tables.append((pending_title, header_row, data_rows))
        pending_title = None
    return tables


def _trim_empty_cols(header_row, data_rows):
    """Drop trailing placeholder columns left over when a narrower table
    sits in a sheet whose other tables are wider (rows are padded to the
    sheet's overall column count)."""
    width = len(header_row)
    while width > 0 and header_row[width-1] is None and all(
        len(r) < width or r[width-1] is None for r in data_rows
    ):
        width -= 1
    return header_row[:width], [r[:width] for r in data_rows]


def build_schema_prompt(dfs: dict[str, pd.DataFrame]) -> str:
    """Build schema description for the LLM to understand the data."""
    parts = []
    for name, df in dfs.items():
        parts.append(f"### Sheet: '{name}'")
        parts.append(f"Shape: {df.shape[0]} rows × {df.shape[1]} columns\n")
        parts.append("Columns:")
        for col in df.columns:
            samples = df[col].dropna().head(3).tolist()
            s_str = ", ".join(str(s) for s in samples)
            parts.append(f"  - '{col}' ({df[col].dtype}, {df[col].notna().sum()} non-null) — e.g. {s_str}")
        parts.append(f"\nFirst 5 rows:\n{df.head().to_string(index=False)}")
        numeric = df.select_dtypes(include=[np.number]).columns.tolist()
        if numeric:
            parts.append(f"\nNumeric summary:\n{df[numeric].describe().to_string()}")
        parts.append("")
    return "\n".join(parts)


# ═══════════════════════════════════════════════════════════════════════════
# STAGE 1 — LLM generates pandas code
# ═══════════════════════════════════════════════════════════════════════════

_CODE_SYS = """\
You are an expert Python/pandas code generator. Given a DataFrame schema and
a user question, write Python code that answers the question.

RULES:
1. DataFrames are in `dfs` (dict: sheet name → DataFrame).
   If only one sheet, also available as `df`.
2. Code MUST end with print() showing the final answer clearly.
3. Use pandas for ALL math — never hardcode numbers.
4. Handle edge cases: check columns exist, handle NaN.
5. Return ONLY executable Python code. No markdown fences, no explanation.
6. Available: pandas (pd), numpy (np), datetime, json.
7. Format numbers nicely: f"{value:,.2f}"
8. IMPORTANT: The data may have had merged cells. If filtering by a column value returns fewer rows than expected, the column may need forward-fill: df['col'] = df['col'].ffill()
"""


def _gen_code(schema, question, model, err_ctx=None):
    msg = f"DATA SCHEMA:\n{schema}\n\nQUESTION: {question}"
    if err_ctx:
        msg += f"\n\nPREVIOUS CODE FAILED:\n{err_ctx}\nFix it. Check column names (case-sensitive), handle NaN."
    code = _call_llm(_CODE_SYS, msg, model).strip()
    if code.startswith("```"):
        lines = code.split("\n")
        end = -1 if lines[-1].strip().startswith("```") else len(lines)
        code = "\n".join(lines[1:end])
    return code


# ═══════════════════════════════════════════════════════════════════════════
# STAGE 2 — Execute code safely
# ═══════════════════════════════════════════════════════════════════════════

def _execute_code(code: str, dfs: dict[str, pd.DataFrame]) -> tuple[bool, str]:
    from contextlib import redirect_stdout, redirect_stderr
    stdout, stderr = io.StringIO(), io.StringIO()
    safe = {
        "__builtins__": {
            "print":print,"len":len,"range":range,"enumerate":enumerate,
            "zip":zip,"sorted":sorted,"reversed":reversed,
            "min":min,"max":max,"sum":sum,"abs":abs,"round":round,
            "int":int,"float":float,"str":str,"bool":bool,
            "list":list,"dict":dict,"tuple":tuple,"set":set,
            "isinstance":isinstance,"type":type,"format":format,
            "map":map,"filter":filter,"any":any,"all":all,
            "ValueError":ValueError,"TypeError":TypeError,
            "KeyError":KeyError,"IndexError":IndexError,"Exception":Exception,
        },
        "pd":pd,"np":np,"dfs":dfs,
        "datetime":__import__("datetime"),"json":__import__("json"),
    }
    if len(dfs) == 1:
        safe["df"] = list(dfs.values())[0]
    try:
        with redirect_stdout(stdout), redirect_stderr(stderr):
            exec(code, safe)
        out = stdout.getvalue().strip()
        err = stderr.getvalue().strip()
        if not out and err: return False, f"No output. Stderr: {err}"
        if not out: return False, "No output — ensure code ends with print()."
        return True, out
    except Exception as e:
        return False, f"{type(e).__name__}: {e}\n{traceback.format_exc()}"


# ═══════════════════════════════════════════════════════════════════════════
# STAGE 3 — LLM formats raw output → readable text
# ═══════════════════════════════════════════════════════════════════════════

_FMT_SYS = """\
Format this data-analysis result into clear, readable text.
- Use natural language, not raw code output.
- Keep EXACT numbers from the computation — never round unless asked.
- Use tables/bullets as appropriate.
- NEVER say "the code produced" — answer as if you computed it.
- If user asked for JSON, return valid JSON.
"""


def _format_answer(question, raw, model):
    return _call_llm(_FMT_SYS, f"Question: {question}\n\nResult:\n{raw}\n\nFormat clearly.", model)


# ═══════════════════════════════════════════════════════════════════════════
# MAIN Q&A — the hybrid pipeline
# ═══════════════════════════════════════════════════════════════════════════

def ask_question(dfs, schema, question, excel_text=None, model=None, **kw):
    """LLM → code → execute → format. Returns dict with answer/code/method."""
    last_err = None
    for attempt in range(MAX_RETRIES + 1):
        try:
            code = _gen_code(schema, question, model, last_err if attempt > 0 else None)
            ok, output = _execute_code(code, dfs)
            if ok:
                answer = _format_answer(question, output, model)
                return {"answer":answer,"code":code,"raw_output":output,"method":"code_execution"}
            last_err = f"Code:\n{code}\n\nError:\n{output}"
            logger.warning("Attempt %d failed: %s", attempt+1, output[:200])
        except Exception as e:
            last_err = str(e)
            logger.warning("Attempt %d error: %s", attempt+1, e)

    # Fallback: direct LLM
    logger.info("Falling back to direct LLM")
    try:
        ctx = excel_text or schema
        ans = _call_llm(
            "You are a data analyst. Answer accurately with step-by-step calculations. Clean readable text.",
            f"DATA:\n{ctx}\n\nQUESTION: {question}", model,
        )
        return {"answer":ans,"code":f"# Fallback (code failed {MAX_RETRIES+1}x)\n# Error: {last_err}",
                "raw_output":"(LLM-computed — verify numbers)","method":"direct_llm"}
    except Exception as e:
        return {"answer":f"Error: {e}","code":"","raw_output":"","method":"error"}


# ═══════════════════════════════════════════════════════════════════════════
# CONVENIENCE
# ═══════════════════════════════════════════════════════════════════════════

def summarize_excel(dfs, schema, model=None):
    return ask_question(dfs, schema,
        "Comprehensive summary: what the data is about, key columns, row counts, "
        "notable patterns, data quality, 3-5 insights with exact numbers.", model=model)


def convert_json_to_text_llm(json_input, model=None):
    try:
        return _call_llm("Convert JSON → clean readable text. Preserve all data. No code blocks.",
                         json_input, model)
    except Exception as e:
        return f"Error: {e}"


def convert_text_to_json_llm(text_input, model=None):
    try:
        raw = _call_llm("Convert text → valid JSON. Return ONLY JSON, no fences, no explanation.",
                        text_input, model).strip()
        return json.dumps(json.loads(raw), indent=2, ensure_ascii=False)
    except json.JSONDecodeError:
        return raw
    except Exception as e:
        return f"Error: {e}"


# ═══════════════════════════════════════════════════════════════════════════
# HELPERS
# ═══════════════════════════════════════════════════════════════════════════

def _find_header_row(rows):
    def _u(r): return set(str(v) for v in r if v is not None)
    for i, row in enumerate(rows):
        nn = [v for v in row if v is not None]
        if not nn: continue
        if len(nn) <= 1 and i+1 < len(rows) and len([v for v in rows[i+1] if v is not None]) > 1:
            continue  # single-cell title row, a richer row follows
        u = _u(row)
        if len(u) <= 1 and len(nn) > 1: continue
        if len(u) < len(nn) and i+1 < len(rows):
            if len(_u(rows[i+1])) > len(u): continue
        return i
    return 0

def _dedup(headers):
    seen, out = {}, []
    for h in headers:
        if h in seen: seen[h] += 1; out.append(f"{h}_{seen[h]}")
        else: seen[h] = 0; out.append(h)
    return out
