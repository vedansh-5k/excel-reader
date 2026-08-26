"""
excel_parser.py — Robust Excel/CSV parser with merged-cell handling.

Handles: merged cells (row + col spans), hierarchical headers,
multiple sheets, mixed data types, large files (smart truncation).
"""

import io
import json
import logging
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

logger = logging.getLogger(__name__)

MAX_ROWS_FULL = 800
MAX_ROWS_SAMPLE = 200
MAX_ROWS_TAIL = 20
MAX_CELL_LENGTH = 500
MAX_CONTEXT_CHARS = 180_000


# ═══════════════════════════════════════════════════════════════════════════
# PUBLIC API
# ═══════════════════════════════════════════════════════════════════════════

def parse_excel(file_path_or_buffer, file_name: str = "upload.xlsx") -> dict:
    ext = Path(file_name).suffix.lower()
    if ext in (".csv", ".tsv"):
        return _parse_csv(file_path_or_buffer, file_name, ext)
    return _parse_xlsx(file_path_or_buffer, file_name)


def sheets_to_text(parsed: dict, sheet_name: str | None = None) -> str:
    parts: list[str] = [f"=== FILE: {parsed['file_name']} ===\n"]
    sheets = parsed["sheets"]
    if sheet_name:
        sheets = [s for s in sheets if s["name"].lower() == sheet_name.lower()]
        if not sheets:
            avail = ", ".join(s["name"] for s in parsed["sheets"])
            return f"Sheet '{sheet_name}' not found. Available: {avail}"

    for sh in sheets:
        parts.append(f"\n--- Sheet: {sh['name']} ---")
        if sh.get("title"):
            parts.append(f"Title: {sh['title']}")
        parts.append(f"Dimensions: {sh['dimensions']['rows']} rows × {sh['dimensions']['cols']} columns")
        if sh.get("merged_regions"):
            shown = sh["merged_regions"][:20]
            parts.append(f"Merged regions: {', '.join(shown)}")
            if len(sh["merged_regions"]) > 20:
                parts.append(f"  ... and {len(sh['merged_regions']) - 20} more")
        if sh.get("headers"):
            parts.append("Headers: " + " | ".join(str(h) for h in sh["headers"]))
        if sh.get("data_types"):
            parts.append(f"Column types: {json.dumps(sh['data_types'], default=str)}")
        if sh.get("summary_stats"):
            parts.append(f"\nSummary statistics:\n{json.dumps(sh['summary_stats'], indent=2, default=str)}")
        if sh["is_truncated"]:
            parts.append(f"\n⚠ Showing {len(sh['rows'])} of {sh['total_rows']} rows (truncated).\n")
        if sh["rows"]:
            hdrs = sh["headers"] or [f"Col{i+1}" for i in range(sh["dimensions"]["cols"])]
            parts.append(_build_md_table(hdrs, sh["rows"]))

    text = "\n".join(parts)
    if len(text) > MAX_CONTEXT_CHARS:
        text = text[:MAX_CONTEXT_CHARS] + "\n\n[... truncated ...]"
    return text


def parsed_to_json(parsed: dict) -> str:
    return json.dumps(parsed, indent=2, default=str, ensure_ascii=False)


def json_to_text(json_str: str) -> str:
    try:
        data = json.loads(json_str)
    except json.JSONDecodeError:
        return f"Invalid JSON:\n{json_str}"
    return _json_val_to_text(data)


def text_to_json(text: str) -> str:
    lines = [ln.strip() for ln in text.strip().splitlines() if ln.strip()]
    result: dict[str, Any] = {}
    cur = None
    for line in lines:
        if ":" in line and not line.startswith("|"):
            k, _, v = line.partition(":")
            k, v = k.strip(), v.strip()
            if v:
                for cast in (int, float):
                    try: v = cast(v); break
                    except ValueError: pass
            result[k] = v; cur = k
        elif line.startswith("|"):
            cells = [c.strip() for c in line.strip("|").split("|")]
            result.setdefault("_table_rows", []).append(cells)
        elif cur and isinstance(result.get(cur), str):
            result[cur] += " " + line
    return json.dumps(result, indent=2, default=str, ensure_ascii=False)


# ═══════════════════════════════════════════════════════════════════════════
# INTERNAL — XLSX
# ═══════════════════════════════════════════════════════════════════════════

def _parse_xlsx(fp, fn):
    wb = load_workbook(fp, data_only=True, read_only=False)
    res = {"file_name": fn, "sheets": []}
    for name in wb.sheetnames:
        res["sheets"].extend(_process_ws(wb[name], name))
    wb.close()
    return res


def _split_blocks(all_rows):
    """Split sheet rows into blocks separated by fully-blank rows."""
    blocks, cur = [], []
    for row in all_rows:
        if all(v is None for v in row):
            if cur:
                blocks.append(cur)
            cur = []
        else:
            cur.append(row)
    if cur:
        blocks.append(cur)
    return blocks


def _trim_block_cols(block):
    """Drop trailing all-blank columns left over when a narrower table sits
    in a sheet whose other tables are wider (rows are padded to the sheet's
    overall column count)."""
    if not block:
        return block
    width = max(len(r) for r in block)
    while width > 0 and all(len(r) < width or r[width-1] is None for r in block):
        width -= 1
    return [r[:width] for r in block]


def _process_ws(ws, ws_name):
    """Process one worksheet into one-or-more table entries. A sheet holding
    several stacked tables (separated by blank rows) yields one entry per
    table so unrelated tables are never merged into a single garbled table."""
    merged_regions = [str(mr) for mr in ws.merged_cells.ranges]

    vmap: dict[tuple[int,int], Any] = {}
    for mr in list(ws.merged_cells.ranges):
        anchor = ws.cell(mr.min_row, mr.min_col).value
        for r in range(mr.min_row, mr.max_row+1):
            for c in range(mr.min_col, mr.max_col+1):
                vmap[(r,c)] = anchor

    mr_count = ws.max_row or 0
    mc_count = ws.max_column or 0
    if mr_count == 0 or mc_count == 0:
        return [_empty(ws_name, merged_regions)]

    all_rows = []
    for r in range(1, mr_count+1):
        row = []
        for c in range(1, mc_count+1):
            v = vmap.get((r,c), ws.cell(r,c).value)
            if isinstance(v, str) and len(v) > MAX_CELL_LENGTH:
                v = v[:MAX_CELL_LENGTH] + "..."
            row.append(v)
        all_rows.append(row)

    blocks = _split_blocks(all_rows)
    if not blocks:
        return [_empty(ws_name, merged_regions)]

    tables, pending_title = [], None
    for block in blocks:
        block = _trim_block_cols(block)
        headers, h_idx, titles = _detect_headers(block, len(block[0]) if block else mc_count)
        data_rows = block[h_idx+1:]
        if not data_rows:
            text = " > ".join(titles) if titles else " ".join(str(v) for v in block[-1] if v is not None)
            pending_title = text or pending_title
            continue
        tables.append((pending_title, headers, data_rows, titles))
        pending_title = None

    if not tables:
        return [_empty(ws_name, merged_regions)]

    sheets, used_names = [], set()
    for idx, (title, headers, data_rows, titles) in enumerate(tables):
        if len(tables) == 1:
            name = ws_name
        elif title:
            short = title.split("|")[0].strip()[:60]
            name = f"{ws_name} - {short}" if short else f"{ws_name} ({idx+1})"
        else:
            name = f"{ws_name} ({idx+1})"
        base, i = name, 2
        while name in used_names:
            name = f"{base} ({i})"; i += 1
        used_names.add(name)

        title_note = " > ".join(titles) if titles else title

        dtypes, samples = {}, {}
        for ci, h in enumerate(headers):
            vals = [r[ci] for r in data_rows if ci < len(r) and r[ci] is not None]
            if vals:
                dtypes[h] = ", ".join(sorted(set(type(v).__name__ for v in vals[:50])))
                samples[h] = str(vals[0])

        total = len(data_rows)
        trunc = total > MAX_ROWS_FULL
        if trunc:
            rows_out = data_rows[:MAX_ROWS_SAMPLE] + [["..."]*len(headers)] + data_rows[-MAX_ROWS_TAIL:]
        else:
            rows_out = data_rows

        sheets.append({
            "name": name, "title": title_note,
            "dimensions": {"rows": total, "cols": len(headers)},
            "headers": headers, "data_types": dtypes, "sample_values": samples,
            "merged_regions": merged_regions,
            "rows": [[_safe(v) for v in row] for row in rows_out],
            "is_truncated": trunc, "total_rows": total,
            "summary_stats": _stats(headers, data_rows),
        })
    return sheets


def _detect_headers(all_rows, max_col):
    def _u(row): return set(str(v) for v in row if v is not None)
    headers, h_idx, titles = [], 0, []
    for i, row in enumerate(all_rows):
        nn = [v for v in row if v is not None]
        if not nn: continue
        if len(nn) <= 1 and i+1 < len(all_rows) and len([v for v in all_rows[i+1] if v is not None]) > 1:
            titles.append(str(nn[0])); continue
        u = _u(row)
        if len(u) <= 1 and len(nn) > 1:
            titles.append(str(nn[0])); continue
        if len(u) < len(nn) and i+1 < len(all_rows):
            if len(_u(all_rows[i+1])) > len(u):
                titles.append(" | ".join(sorted(u))); continue
        headers = [str(v) if v is not None else f"Col{j+1}" for j,v in enumerate(row)]
        h_idx = i; break
    if not headers:
        for i, row in enumerate(all_rows):
            if any(v is not None for v in row):
                headers = [str(v) if v is not None else f"Col{j+1}" for j,v in enumerate(row)]
                h_idx = i; break
        if not headers:
            headers = [f"Col{j+1}" for j in range(max_col)]
    return headers, h_idx, titles


# ═══════════════════════════════════════════════════════════════════════════
# INTERNAL — CSV
# ═══════════════════════════════════════════════════════════════════════════

def _parse_csv(fp, fn, ext):
    sep = "\t" if ext == ".tsv" else ","
    if isinstance(fp, (str, Path)):
        df = pd.read_csv(fp, sep=sep, dtype=str, on_bad_lines="skip")
    else:
        df = pd.read_csv(io.BytesIO(fp.read()), sep=sep, dtype=str, on_bad_lines="skip")
    hdrs = list(df.columns)
    total = len(df)
    trunc = total > MAX_ROWS_FULL
    show = pd.concat([df.head(MAX_ROWS_SAMPLE), df.tail(MAX_ROWS_TAIL)]) if trunc else df
    clean = [[_safe(v) for v in row] for row in show.values.tolist()]
    df_num = df.apply(pd.to_numeric, errors="coerce")
    st = {}
    for c in df_num.columns:
        if df_num[c].notna().sum() > 0:
            st[c] = {"min":float(df_num[c].min()),"max":float(df_num[c].max()),
                      "mean":round(float(df_num[c].mean()),2),"count":int(df_num[c].notna().sum())}
    return {"file_name":fn,"sheets":[{
        "name":"Sheet1","title":None,"dimensions":{"rows":total,"cols":len(hdrs)},
        "headers":hdrs,"data_types":{h:"str" for h in hdrs},
        "sample_values":{h:str(df[h].iloc[0]) if len(df)>0 else "" for h in hdrs},
        "merged_regions":[],"rows":clean,"is_truncated":trunc,"total_rows":total,
        "summary_stats":st or None,
    }]}


# ═══════════════════════════════════════════════════════════════════════════
# INTERNAL — helpers
# ═══════════════════════════════════════════════════════════════════════════

def _empty(name, merged):
    return {"name":name,"title":None,"dimensions":{"rows":0,"cols":0},
            "headers":[],"data_types":{},"sample_values":{},"merged_regions":merged,
            "rows":[],"is_truncated":False,"total_rows":0,"summary_stats":None}

def _stats(headers, data_rows):
    st = {}
    for ci, h in enumerate(headers):
        nums = []
        for r in data_rows:
            if ci < len(r):
                v = r[ci]
                if isinstance(v,(int,float)) and not isinstance(v,bool): nums.append(float(v))
                elif isinstance(v,str):
                    try: nums.append(float(v.replace(",","")))
                    except: pass
        if nums:
            st[h] = {"min":min(nums),"max":max(nums),"mean":round(sum(nums)/len(nums),2),
                      "sum":round(sum(nums),2),"count":len(nums)}
    return st or None

def _build_md_table(headers, rows):
    w = [len(str(h)) for h in headers]
    for row in rows[:5]:
        for i,v in enumerate(row):
            if i < len(w): w[i] = max(w[i], min(len(str(v)),30))
    hdr = "| " + " | ".join(str(h)[:30].ljust(w[i]) for i,h in enumerate(headers)) + " |"
    sep = "| " + " | ".join("-"*min(x,30) for x in w) + " |"
    lines = [hdr, sep]
    for row in rows:
        cells = []
        for i in range(len(headers)):
            v = row[i] if i < len(row) else ""
            ww = w[i] if i < len(w) else 10
            cells.append(str(v)[:30].ljust(ww))
        lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)

def _safe(v):
    if v is None: return ""
    if isinstance(v,float): return str(int(v)) if v==int(v) else f"{v:.4f}".rstrip("0").rstrip(".")
    return str(v)

def _json_val_to_text(obj, indent=0):
    p = "  "*indent
    if isinstance(obj,dict):
        return "\n".join(f"{p}{k}:\n{_json_val_to_text(v,indent+1)}" if isinstance(v,(dict,list))
                         else f"{p}{k}: {v}" for k,v in obj.items())
    if isinstance(obj,list):
        return "\n".join(f"{p}Item {i+1}:\n{_json_val_to_text(x,indent+1)}" if isinstance(x,dict)
                         else f"{p}- {x}" for i,x in enumerate(obj))
    return f"{p}{obj}"
